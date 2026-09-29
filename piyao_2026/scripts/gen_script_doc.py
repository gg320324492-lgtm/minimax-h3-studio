"""Generate the submission script document (Word) for piyao_2026.

Reads shots.json + timeline.json (real timecodes) and writes
09_final/它只是换了一个地名_作品脚本.docx with:
  title block, synopsis, storyboard table (14 shots, real timecodes),
  narration full text, production & originality notes.
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT = Path(r'E:/Minimax-H3/piyao_2026')
CFG = json.loads((PROJECT / '00_project/shots.json').read_text(encoding='utf-8'))
TL = json.loads((PROJECT / '00_project/timeline.json').read_text(encoding='utf-8'))

ORDER = ['S01', 'S02', 'S03', 'S04', 'S05', 'S06', 'S07', 'S08',
         'S09', 'S10', 'S11', 'S12', 'S13', 'S14']

DESC_CN = {
    'S01': '夜晚，大学生宿舍。台灯暖光，书桌上有书本、笔记本电脑。一名穿浅灰色卫衣的大学生在桌前，桌上手机突然亮起冷白光，映亮半边脸。他望向手机，拿起手机。（对应群聊弹窗字卡：“突发！河南某地大片麦田起火！速转！”）',
    'S02': '手机特写：屏幕中正播放一段麦田起火视频；学生拇指缓慢移向屏幕右下角，即将点下转发的瞬间，手指悬停、画面定格。叠字幕“等等。你看到的，未必就是刚刚发生的。”左下角持续角标“情景演绎｜AI技术辅助生成”。',
    'S03': '这段麦田起火视频被“抽离”出来，化作黑色数字空间中一块悬浮的发光屏幕，微微旋转漂浮，四周有细碎的数字化粒子与光线，克制、高级的信息空间质感。',
    'S04': '黑色空间中，同一段起火视频保持画面中央不变；大量空白标签贴纸接连飞来贴上又被撕下，红色感叹号徽章与定位图标不断出现又消失——内容没变，包装在被反复更换。',
    'S05': '一块发光屏幕分裂成许多悬浮手机屏：宿舍里的学生、客厅里的家长、戴老花镜的老人、院子里的农户、公交上的年轻人，人人低头看着同一条“突发消息”，冷色手机光，链式扩散感。',
    'S06': '信息污染隐喻镜头：黑色像素颗粒、红色警示徽章与感叹号碎片从无数手机屏中涌出，像雾霾般漫过城市夜空，再向乡村与麦田蔓延。冷色调、红点缀，克制不安。',
    'S07': '回到宿舍：学生拇指悬停却没有按下转发，神情由焦虑转为冷静，微微放下手机，改点搜索。冷光映面，气氛由紧张转为理性。',
    'S08': '蓝白色理性核验空间：左侧悬浮着那段起火视频，旁边出现时间轴与检索卡片，一个更早的日期标记被点亮——同一段视频其实早已存在。（叠大字卡“第一步：查时间”）',
    'S09': '核验继续：放大框逐一检视视频内的地貌、远山、天空等细节；一个错误的红色定位标记碎裂消散，四周排开比对面板。（叠大字卡“第二步：查地点”）',
    'S10': '明亮洁净的数字空间：一排信息卡悬浮而来，灰色的无出处卡片渐暗退场，中央一张卡片被蓝色盾形光晕环绕，稳定清晰——权威来源与非权威来源的对比。（叠大字卡“第三步：查来源”）',
    'S11': '三张洁净的发光卡片并排而立；下方红色警示徽章被三束光照射，碎裂成红色粒子消散。明亮、对称、克制的收束感。（叠字卡“查时间｜查地点｜查来源”）',
    'S12': '清晨的河南麦田：金色阳光铺满起伏的麦浪，薄雾通透，远处收割机缓缓行驶，乡间小路整洁。航拍低空缓推，真实、安宁、有力量。',
    'S13': '温暖群像：学生骑车经过麦田边的乡村公路，田埂上农户驻足查看麦穗。柔光、暖色、平静真实——清朗的网络空间与真实生活息息相关。',
    'S14': '清晨宿舍：学生双手持机，拇指从容点开搜索，屏幕泛起柔和白光；他微微一笑，把手机轻轻放回桌面的书旁。明亮笃定，与开头呼应。',
}

NARR_OF = {n['id']: n for n in CFG['narration'] if not n.get('skip')}
SUBS = TL['subtitle_events']


def tc(t):
    m, s = int(t) // 60, t - int(t) // 60 * 60
    return f'{m:02d}:{s:05.2f}'


def set_font(run, size=10.5, bold=False, color=None, name='微软雅黑'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text, size=10.5, bold=False, align=None, color=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size, bold, color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def main():
    doc = Document()

    # page + default style font
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2.2)
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # ---- title block ----
    para(doc, '2026年度互联网辟谣优秀作品征集活动', 12, True, WD_ALIGN_PARAGRAPH.CENTER, (96, 96, 96))
    para(doc, '音视频类参赛作品脚本', 14, True, WD_ALIGN_PARAGRAPH.CENTER, (96, 96, 96))
    para(doc, '《它只是换了一个地名》', 24, True, WD_ALIGN_PARAGRAPH.CENTER, (20, 30, 50), 10)
    para(doc, '活动主题：“豫”你一起 同心护网', 12, False, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, '作品类型：公益宣传短片（音视频类）｜时长：约1分49秒｜画幅：16:9横屏 1920×1080',
         10.5, False, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, '本内容经AI技术辅助生成', 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, (140, 30, 30), 14)

    # ---- synopsis ----
    doc.add_heading('一、作品简介', level=1)
    synopsis = ('《它只是换了一个地名》是一部面向公众的网络辟谣公益短片。作品聚焦“旧视频、旧图片被重新包装、'
                '旧闻新传”这一常见谣言套路：同一段麦田视频，被换上“刚刚发生”“河南某地”的时间、地点与标题后，'
                '便可能摇身变成一条新谣言。短片以大学生夜间刷手机的生活化场景切入，用视觉化手法呈现谣言的传播'
                '机制与“信息污染”，并明确给出“查时间、查地点、查来源”三步识谣方法，倡导网民转发前多一步核实。'
                '作品融入河南麦田、乡村公路等本地意象，以“转发前，给真相十秒”收束，兼具科普性、时效性与传播性。')
    para(doc, synopsis, 10.5)
    para(doc, '核心传播句：转发前，给真相十秒。', 10.5, True)

    # ---- storyboard table ----
    doc.add_heading('二、分镜脚本（成片时码）', level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['镜号/时码', '画面内容（场景描述·分镜设计）', '旁白台词', '屏幕字幕/字卡', '音效', '音乐']
    widths = [Cm(2.2), Cm(6.4), Cm(3.6), Cm(2.8), Cm(1.8), Cm(1.6)]
    for i, (h, w) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = w
        r = cell.paragraphs[0].add_run(h)
        set_font(r, 9.5, True)

    sfx_map = {}
    for e in TL['sfx_events']:
        sfx_map.setdefault(e['start'], []).append(e['file'].split('/')[-1].replace('SFX_', '').replace('.m4a', ''))

    for idx, sid in enumerate(ORDER, 1):
        b = TL['shot_boundaries'][sid]
        row = table.add_row()
        cells = row.cells
        for c, w in zip(cells, widths):
            c.width = w
        r0 = cells[0].paragraphs[0]
        set_font(r0.add_run(f'{idx:02d} {sid}\n'), 9, True)
        set_font(r0.add_run(f'{tc(b["start"])}–{tc(b["end"])}\n({b["duration"]:.1f}s)'), 8.5, False, (110, 110, 110))
        set_font(cells[1].paragraphs[0].add_run(DESC_CN[sid]), 9, False)
        narr = [n for n in CFG['narration'] if n.get('shot') == sid and not n.get('skip')]
        set_font(cells[2].paragraphs[0].add_run(narr[0]['text'] if narr else '（无旁白）'), 9)
        subs = [e['text'] for e in SUBS
                if b['start'] - 0.01 <= e['start'] < b['end'] and e['style'] not in ('corner',)]
        set_font(cells[3].paragraphs[0].add_run('\n'.join(subs) if subs else '—'), 8.5)
        sfx = sorted({s for t, names in sfx_map.items() if b['start'] - 0.5 <= t < b['end'] for s in names})
        set_font(cells[4].paragraphs[0].add_run('\n'.join(sfx) if sfx else '—'), 8.5)
        phase = ('紧张低音' if sid in ('S01', 'S02', 'S03') else
                 '焦虑脉冲' if sid in ('S04', 'S05', 'S06') else
                 '理性清亮' if sid in ('S07', 'S08', 'S09', 'S10', 'S11') else '温暖弦乐')
        set_font(cells[5].paragraphs[0].add_run(phase), 8.5)

    # end card row
    row = table.add_row()
    for c, w in zip(row.cells, widths):
        c.width = w
    e0 = TL['endcard_start']
    set_font(row.cells[0].paragraphs[0].add_run(f'片尾\n{tc(e0)}–{tc(TL["total_duration"])}'), 9, True)
    set_font(row.cells[1].paragraphs[0].add_run('片尾定版卡：深蓝渐变底，主文案与标识依次呈现'), 9)
    set_font(row.cells[2].paragraphs[0].add_run('（旁白已结束）'), 9)
    set_font(row.cells[3].paragraphs[0].add_run(
        '转发前，给真相十秒；“豫”你一起 同心护网；《它只是换了一个地名》·2026年度互联网辟谣优秀作品征集；'
        '本内容经AI技术辅助生成；本片为公益宣传情景演绎，非真实灾情记录'), 8.5)
    set_font(row.cells[4].paragraphs[0].add_run('END_DING'), 8.5)
    set_font(row.cells[5].paragraphs[0].add_run('温暖收束'), 8.5)

    # ---- narration full text ----
    doc.add_heading('三、旁白全文', level=1)
    for n in CFG['narration']:
        if n.get('skip'):
            continue
        para(doc, n['text'], 10.5, False, space_after=3)
    para(doc, '配音说明：AI语音合成（Kokoro TTS，男声），语速随段落微调。', 9, False, color=(110, 110, 110))

    # ---- production notes ----
    doc.add_heading('四、制作说明与原创声明', level=1)
    for t in [
        '1. 本作品为2025年12月至2026年9月期间原创制作的公益宣传短片，未侵犯任何第三方权益；画面、旁白、配乐、音效均为原创或程序化合成，无第三方版权素材。',
        '2. 画面由开源视频生成模型（MiniMax H3，本地部署）辅助生成，剪辑、配音、配乐、字幕、混音由作者完成；全片按规范添加“本内容经AI技术辅助生成”标识（画面左下角持续角标及片尾卡注明）。',
        '3. 片中“麦田起火”为公益宣传情景演绎画面，非真实灾情记录；片中群聊弹窗、检索界面等文字信息均为后期制作叠加的示意内容。',
        '4. 作品紧扣“‘豫’你一起 同心护网”主题，聚焦生活科普辟谣与防谣技巧普及，结合河南麦田、乡村等本地元素创作。',
    ]:
        para(doc, t, 10.5)

    out = PROJECT / '09_final' / '它只是换了一个地名_作品脚本.docx'
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f'saved: {out}')


if __name__ == '__main__':
    main()
